from __future__ import annotations

import io

import openpyxl
import pytest

from export import build_excel_report, export_filename


@pytest.fixture
def workbook(report):
    return openpyxl.load_workbook(io.BytesIO(build_excel_report(report)))


def test_workbook_has_all_sheets(workbook):
    assert workbook.sheetnames == ["Summary", "Totals", "Ratios", "Income", "Expenses", "Alerts"]


def test_summary_key_figures_are_real_numbers(workbook):
    ws = workbook["Summary"]
    assert ws["A1"].value == "Hotel Jan vs Feb 2026"
    assert "January 2026 → February 2026" in ws["A2"].value
    assert ws["A5"].value == "Total income"
    assert ws["B5"].value == 63_000_000 and ws["C5"].value == 54_500_000
    assert ws["D5"].value == -8_500_000
    assert ws["E5"].value == pytest.approx(-13.49, abs=0.01)
    assert ws["A8"].value == "Profit margin"
    assert ws["C8"].value == pytest.approx(20.0, abs=0.01)


def test_summary_keeps_myanmar_text_and_strips_markdown(workbook):
    text = "\n".join(str(c.value) for row in workbook["Summary"].iter_rows() for c in row if c.value)
    assert "ဝင်ငွေ" in text
    assert "**" not in text
    assert "• 🔴 Maintenance +209%" in text


def test_category_sheet_has_rows_and_sum_formula(workbook, report):
    ws = workbook["Expenses"]
    n = len(report["metrics"]["expense_by_category"])
    assert ws.cell(1, 1).value == "Category"
    categories = {ws.cell(r, 1).value for r in range(2, n + 2)}
    assert "Renovation" in categories and "Maintenance" in categories
    assert ws.cell(n + 2, 1).value == "Total"
    assert ws.cell(n + 2, 3).value == f"=SUM(C2:C{n + 1})"


def test_alerts_sheet_lists_every_alert(workbook, report):
    ws = workbook["Alerts"]
    messages = [ws.cell(r, 2).value for r in range(2, ws.max_row + 1)]
    assert messages == [a["message"] for a in report["alerts"]]
    assert ws.cell(2, 1).value in {"Critical", "Warning", "Note"}


def test_export_survives_missing_summary_and_alerts(report):
    report["summary"] = None
    report["alerts"] = []
    wb = openpyxl.load_workbook(io.BytesIO(build_excel_report(report)))
    assert wb["Alerts"].max_row == 1


def test_export_filename_is_safe(report):
    report["title"] = "Hotel: Jan/Feb 2026 ✓"
    name = export_filename(report)
    assert name.endswith("_1.xlsx")
    assert "/" not in name and ":" not in name and " " not in name
