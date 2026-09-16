"""
Excel export of a comparison report, for sharing with the board or the accountant.

build_excel_report(rec) -> bytes of an .xlsx workbook with these sheets:
    Summary · Totals · Ratios · Income · Expenses · Alerts
Numbers are written as real numbers with number formats (not text), so the recipient
can keep calculating on them.
"""

from __future__ import annotations

import io
import re
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from finance import display_period, is_missing, safe_div
from i18n import EN, alert_text, item_name, ratio_name, t

NAVY = "0F2A4A"
HEADER_FILL = PatternFill("solid", fgColor=NAVY)
HEADER_FONT = Font(bold=True, color="FFFFFF")
TITLE_FONT = Font(bold=True, size=16, color=NAVY)
LABEL_FONT = Font(bold=True, color="555555")
GOOD = Font(color="1A7F4B", bold=True)
BAD = Font(color="C0392B", bold=True)
THIN = Side(style="thin", color="DDDDDD")
BOX = Border(bottom=THIN)
SEVERITY_FILL = {
    "error": PatternFill("solid", fgColor="FDECEA"),
    "warning": PatternFill("solid", fgColor="FFF6E0"),
    "info": PatternFill("solid", fgColor="EAF2FD"),
}
MONEY = "#,##0"
SIGNED_MONEY = "+#,##0;-#,##0;0"
PCT = '+0.0"%";-0.0"%";0.0"%"'
PLAIN_PCT = '0.0"%"'
MULT = '0.00"x"'


def _num(v: Any) -> float | None:
    return None if is_missing(v) else float(v)


def _header(ws, row: int, values: list[str]) -> None:
    for col, v in enumerate(values, start=1):
        c = ws.cell(row=row, column=col, value=v)
        c.fill, c.font = HEADER_FILL, HEADER_FONT
        c.alignment = Alignment(vertical="center")


def _autosize(ws, min_width: int = 10, max_width: int = 60) -> None:
    widths: dict[int, int] = {}
    for row in ws.iter_rows():
        for c in row:
            if c.value is None:
                continue
            longest = max(len(line) for line in str(c.value).splitlines() or [""])
            widths[c.column] = max(widths.get(c.column, 0), longest)
    for col, w in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = max(min_width, min(max_width, w + 2))


def _colour_change(cell, value: float | None, higher_is_good: bool = True) -> None:
    if value is None or value == 0:
        return
    cell.font = GOOD if (value > 0) == higher_is_good else BAD


def _strip_markdown(text: str) -> str:
    text = re.sub(r"\*\*(.+?)\*\*", r"\1", text or "")
    return re.sub(r"^\s*[-*]\s+", "• ", text, flags=re.MULTILINE)


def build_excel_report(rec: dict[str, Any], lang: str = EN) -> bytes:
    d1, d2 = rec["data1"], rec["data2"]
    m = rec.get("metrics") or {}
    l1, l2 = rec.get("label1") or "Previous", rec.get("label2") or "Current"
    cur = rec.get("currency") or ""
    wb = Workbook()

    # ---- Summary -----------------------------------------------------------
    ws = wb.active
    ws.title = "Summary"
    ws["A1"] = rec.get("title") or "Financial comparison"
    ws["A1"].font = TITLE_FONT
    ws["A2"] = (f"{display_period(rec.get('period1'), l1)} → {display_period(rec.get('period2'), l2)}"
                f"   ·   {t('xl_currency', lang)}: {cur or '-'}   ·   {t('xl_generated', lang)}")
    ws["A2"].font = LABEL_FONT

    _header(ws, 4, [t("xl_key_figure", lang), l1, l2, t("col_change", lang), t("col_change_pct", lang)])
    rows = [(t("xl_total_income", lang), "total_income", True),
            (t("xl_total_expense", lang), "total_expense", False),
            (t("xl_net_profit", lang), "net_profit", True)]
    r = 5
    for name, key, good_up in rows:
        a, b = _num(d1.get(key)), _num(d2.get(key))
        ws.cell(r, 1, name).font = LABEL_FONT
        ws.cell(r, 2, a).number_format = MONEY
        ws.cell(r, 3, b).number_format = MONEY
        diff = None if a is None or b is None else b - a
        c = ws.cell(r, 4, diff)
        c.number_format = SIGNED_MONEY
        _colour_change(c, diff, good_up)
        pct = None if not a else (b - a) / abs(a) * 100
        c = ws.cell(r, 5, pct)
        c.number_format = PCT
        _colour_change(c, pct, good_up)
        r += 1
    pm1 = safe_div(d1["net_profit"], d1["total_income"])
    pm2 = safe_div(d2["net_profit"], d2["total_income"])
    ws.cell(r, 1, t("xl_profit_margin", lang)).font = LABEL_FONT
    ws.cell(r, 2, None if pm1 is None else pm1 * 100).number_format = PLAIN_PCT
    ws.cell(r, 3, None if pm2 is None else pm2 * 100).number_format = PLAIN_PCT
    if pm1 is not None and pm2 is not None:
        c = ws.cell(r, 4, (pm2 - pm1) * 100)
        c.number_format = '+0.0" pts";-0.0" pts"'
        _colour_change(c, pm2 - pm1)
    for row in ws.iter_rows(min_row=5, max_row=r):
        for c in row:
            c.border = BOX

    alerts = rec.get("alerts") or []
    counts = {s: sum(1 for a in alerts if a.get("severity") == s) for s in ("error", "warning", "info")}
    r += 2
    ws.cell(r, 1, t("xl_alerts", lang)).font = LABEL_FONT
    ws.cell(r, 2, t("xl_alert_counts", lang, e=counts["error"], w=counts["warning"], i=counts["info"]))

    r += 2
    ws.cell(r, 1, t("xl_summary", lang)).font = TITLE_FONT
    r += 1
    for line in _strip_markdown(rec.get("summary") or t("xl_no_summary", lang)).splitlines():
        if line.strip():
            c = ws.cell(r, 1, line.strip())
            c.alignment = Alignment(wrap_text=True, vertical="top")
            ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=5)
            ws.row_dimensions[r].height = max(15, 15 * (len(line) // 90 + 1))
            r += 1
    ws.column_dimensions["A"].width = 26
    for col in "BCDE":
        ws.column_dimensions[col].width = 18

    # ---- Totals ------------------------------------------------------------
    ws = wb.create_sheet("Totals")
    _header(ws, 1, [t("col_item", lang), l1, l2, t("col_variance", lang), t("col_combined", lang),
                    t("col_change_pct", lang)])
    for i, tot in enumerate(m.get("totals") or [], start=2):
        ws.cell(i, 1, item_name(tot["Item"], lang))
        ws.cell(i, 2, _num(tot.get(l1))).number_format = MONEY
        ws.cell(i, 3, _num(tot.get(l2))).number_format = MONEY
        ws.cell(i, 4, _num(tot.get("Net Variance"))).number_format = SIGNED_MONEY
        ws.cell(i, 5, _num(tot.get("Combined"))).number_format = MONEY
        ws.cell(i, 6, _num(tot.get("Change %"))).number_format = PCT
    _autosize(ws)

    # ---- Ratios ------------------------------------------------------------
    ws = wb.create_sheet("Ratios")
    _header(ws, 1, [t("col_metric", lang), l1, l2])
    for i, rt in enumerate(m.get("ratios") or [], start=2):
        fmt = MULT if rt.get("unit") == "x" else PLAIN_PCT
        ws.cell(i, 1, ratio_name(rt["Metric"], lang))
        ws.cell(i, 2, _num(rt.get(l1))).number_format = fmt
        ws.cell(i, 3, _num(rt.get(l2))).number_format = fmt
    _autosize(ws)

    # ---- Categories --------------------------------------------------------
    for sheet, key, good_up in (("Income", "income_by_category", True),
                                ("Expenses", "expense_by_category", False)):
        ws = wb.create_sheet(sheet)
        _header(ws, 1, [t("col_category", lang), l1, l2, t("col_difference", lang),
                        t("col_change_pct", lang)])
        rows_ = m.get(key) or []
        for i, row in enumerate(rows_, start=2):
            ws.cell(i, 1, row["category"])
            ws.cell(i, 2, _num(row.get(l1))).number_format = MONEY
            ws.cell(i, 3, _num(row.get(l2))).number_format = MONEY
            c = ws.cell(i, 4, _num(row.get("Difference")))
            c.number_format = SIGNED_MONEY
            _colour_change(c, _num(row.get("Difference")), good_up)
            ws.cell(i, 5, _num(row.get("Change %"))).number_format = PCT
        if rows_:
            n = len(rows_) + 2
            ws.cell(n, 1, t("xl_total", lang)).font = Font(bold=True)
            for col in (2, 3, 4):
                letter = get_column_letter(col)
                c = ws.cell(n, col, f"=SUM({letter}2:{letter}{n - 1})")
                c.font = Font(bold=True)
                c.number_format = SIGNED_MONEY if col == 4 else MONEY
        _autosize(ws)

    # ---- Alerts ------------------------------------------------------------
    ws = wb.create_sheet("Alerts")
    _header(ws, 1, [t("xl_severity", lang), t("xl_message", lang)])
    for i, a in enumerate(alerts, start=2):
        sev = a.get("severity", "info")
        sev_label = t(f"sev_{sev}", lang) if sev in ("error", "warning", "info") else sev
        ws.cell(i, 1, sev_label).fill = SEVERITY_FILL.get(sev, SEVERITY_FILL["info"])
        c = ws.cell(i, 2, alert_text(a, lang))
        c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions["A"].width = 12
    ws.column_dimensions["B"].width = 100

    for sheet in wb.worksheets:
        sheet.freeze_panes = "A2" if sheet.title != "Summary" else None

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def export_filename(rec: dict[str, Any]) -> str:
    base = re.sub(r"[^A-Za-z0-9._-]+", "_", rec.get("title") or "report").strip("_") or "report"
    return f"{base}_{rec.get('id', 'new')}.xlsx"
